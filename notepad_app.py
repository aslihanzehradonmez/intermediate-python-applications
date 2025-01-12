import tkinter as tk
from tkinter import ttk
from tkinter.scrolledtext import ScrolledText
import threading
import asyncio
import datetime
import os
from tkinter.filedialog import asksaveasfilename
from docx import Document
from fpdf import FPDF

class AdvancedNotepadApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Notepad App")
        self.root.geometry("1000x700")
        
        # Main Frame
        self.main_frame = ttk.Frame(self.root, padding="10")
        self.main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        self.root.columnconfigure(0, weight=1)
        self.root.rowconfigure(0, weight=1)

        # Menu bar
        self.menu_bar = tk.Menu(self.root)
        self.root.config(menu=self.menu_bar)
        
        self.file_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.file_menu.add_command(label="New", command=self.new_file)
        self.file_menu.add_command(label="Open", command=self.open_file)
        self.file_menu.add_command(label="Save As (TXT)", command=self.save_file_as_txt)
        self.file_menu.add_command(label="Save As (DOCX)", command=self.save_file_as_docx)
        self.file_menu.add_command(label="Save As (PDF)", command=self.save_file_as_pdf)
        self.file_menu.add_separator()
        self.file_menu.add_command(label="Exit", command=self.root.quit)
        self.menu_bar.add_cascade(label="File", menu=self.file_menu)

        self.edit_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.edit_menu.add_command(label="Undo", command=self.undo)
        self.edit_menu.add_command(label="Redo", command=self.redo)
        self.edit_menu.add_separator()
        self.edit_menu.add_command(label="Cut", command=self.cut_text)
        self.edit_menu.add_command(label="Copy", command=self.copy_text)
        self.edit_menu.add_command(label="Paste", command=self.paste_text)
        self.menu_bar.add_cascade(label="Edit", menu=self.edit_menu)
        
        # Toolbar
        self.toolbar = ttk.Frame(self.main_frame)
        self.toolbar.grid(row=0, column=0, sticky=(tk.W, tk.E))
        
        self.bold_btn = ttk.Button(self.toolbar, text="Bold", command=self.make_bold)
        self.bold_btn.pack(side=tk.LEFT, padx=2)
        
        self.italic_btn = ttk.Button(self.toolbar, text="Italic", command=self.make_italic)
        self.italic_btn.pack(side=tk.LEFT, padx=2)

        self.search_replace_panel = None
        self.search_replace_btn = ttk.Button(self.toolbar, text="Search/Replace", command=self.toggle_search_replace)
        self.search_replace_btn.pack(side=tk.LEFT, padx=2)
        
        # Text Editor
        self.text_area = ScrolledText(self.main_frame, wrap=tk.WORD, font=("Helvetica", 12))
        self.text_area.grid(row=1, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        self.main_frame.rowconfigure(1, weight=1)
        self.main_frame.columnconfigure(0, weight=1)

        # Notifications Panel
        self.notifications_panel = tk.Text(self.main_frame, height=5, wrap=tk.WORD, state=tk.DISABLED, bg="#f5f5f5")
        self.notifications_panel.grid(row=3, column=0, sticky=(tk.W, tk.E))

        # AZD Label
        self.azd_label = ttk.Label(self.main_frame, text="AZD", anchor=tk.CENTER, font=("Helvetica", 10))
        self.azd_label.grid(row=4, column=0, sticky=(tk.W, tk.E))

        # Variables
        self.current_file = None

    def toggle_search_replace(self):
        if self.search_replace_panel:
            self.search_replace_panel.destroy()
            self.search_replace_panel = None
        else:
            self.search_replace_panel = ttk.Frame(self.main_frame)
            self.search_replace_panel.grid(row=2, column=0, sticky=(tk.W, tk.E))

            ttk.Label(self.search_replace_panel, text="Search:").grid(row=0, column=0, padx=5, pady=5, sticky=tk.W)
            search_entry = ttk.Entry(self.search_replace_panel)
            search_entry.grid(row=0, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

            ttk.Label(self.search_replace_panel, text="Replace:").grid(row=1, column=0, padx=5, pady=5, sticky=tk.W)
            replace_entry = ttk.Entry(self.search_replace_panel)
            replace_entry.grid(row=1, column=1, padx=5, pady=5, sticky=(tk.W, tk.E))

            search_btn = ttk.Button(self.search_replace_panel, text="Search", command=lambda: self.search_text(search_entry.get()))
            search_btn.grid(row=0, column=2, padx=5, pady=5)

            replace_btn = ttk.Button(self.search_replace_panel, text="Replace", command=lambda: self.replace_text(search_entry.get(), replace_entry.get()))
            replace_btn.grid(row=1, column=2, padx=5, pady=5)

    def log_message(self, message):
        """Logs messages to the notifications panel."""
        self.notifications_panel.config(state=tk.NORMAL)
        timestamp = datetime.datetime.now().strftime("[%H:%M:%S] ")
        self.notifications_panel.insert(tk.END, timestamp + message + "\n")
        self.notifications_panel.config(state=tk.DISABLED)
        self.notifications_panel.see(tk.END)

    def new_file(self):
        """Clears the text area for a new file."""
        self.text_area.delete(1.0, tk.END)
        self.current_file = None
        self.log_message("Started a new file.")

    def open_file(self):
        """Opens a file asynchronously."""
        from tkinter.filedialog import askopenfilename

        def load_file(filepath):
            with open(filepath, 'r') as file:
                content = file.read()
                self.text_area.delete(1.0, tk.END)
                self.text_area.insert(tk.END, content)
                self.current_file = filepath
            self.log_message(f"File opened: {os.path.basename(filepath)}")

        async def async_open_file():
            filepath = askopenfilename()
            if filepath:
                load_file(filepath)
        
        threading.Thread(target=lambda: asyncio.run(async_open_file())).start()

    def save_file_as_txt(self):
        """Saves the current content to a new TXT file."""
        filepath = asksaveasfilename(defaultextension=".txt",
                                      filetypes=[("Text files", "*.txt"), ("All files", "*.*")])
        if filepath:
            with open(filepath, 'w') as file:
                file.write(self.text_area.get(1.0, tk.END).strip())
            self.current_file = filepath
            self.log_message(f"File saved as: {os.path.basename(filepath)}")

    def save_file_as_docx(self):
        """Saves the current content to a DOCX file."""
        filepath = asksaveasfilename(defaultextension=".docx",
                                      filetypes=[("Word Documents", "*.docx"), ("All files", "*.*")])
        if filepath:
            document = Document()
            document.add_paragraph(self.text_area.get(1.0, tk.END).strip())
            document.save(filepath)
            self.log_message(f"File saved as: {os.path.basename(filepath)}")

    def save_file_as_pdf(self):
        """Saves the current content to a PDF file."""
        filepath = asksaveasfilename(defaultextension=".pdf",
                                      filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")])
        if filepath:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.set_font("Arial", size=12)
            content = self.text_area.get(1.0, tk.END).strip()
            for line in content.split("\n"):
                pdf.cell(200, 10, txt=line, ln=True)
            pdf.output(filepath)
            self.log_message(f"File saved as: {os.path.basename(filepath)}")

    def undo(self):
        """Undo the last operation."""
        try:
            self.text_area.edit_undo()
            self.log_message("Undo action performed.")
        except tk.TclError:
            self.log_message("Nothing to undo.")

    def redo(self):
        """Redo the last undone operation."""
        try:
            self.text_area.edit_redo()
            self.log_message("Redo action performed.")
        except tk.TclError:
            self.log_message("Nothing to redo.")

    def cut_text(self):
        """Cuts the selected text."""
        try:
            self.text_area.event_generate("<<Cut>>")
            self.log_message("Cut selected text.")
        except tk.TclError:
            self.log_message("No text selected to cut.")

    def copy_text(self):
        """Copies the selected text."""
        try:
            self.text_area.event_generate("<<Copy>>")
            self.log_message("Copied selected text.")
        except tk.TclError:
            self.log_message("No text selected to copy.")

    def paste_text(self):
        """Pastes the clipboard content."""
        try:
            self.text_area.event_generate("<<Paste>>")
            self.log_message("Pasted clipboard content.")
        except tk.TclError:
            self.log_message("Nothing to paste.")

    def make_bold(self):
        """Makes selected text bold."""
        try:
            self.text_area.tag_configure("bold", font=("Helvetica", 12, "bold"))
            current_tags = self.text_area.tag_names(tk.SEL_FIRST)
            if "bold" in current_tags:
                self.text_area.tag_remove("bold", tk.SEL_FIRST, tk.SEL_LAST)
            else:
                self.text_area.tag_add("bold", tk.SEL_FIRST, tk.SEL_LAST)
            self.log_message("Toggled bold on selected text.")
        except tk.TclError:
            self.log_message("No text selected to apply bold.")

    def make_italic(self):
        """Makes selected text italic."""
        try:
            self.text_area.tag_configure("italic", font=("Helvetica", 12, "italic"))
            current_tags = self.text_area.tag_names(tk.SEL_FIRST)
            if "italic" in current_tags:
                self.text_area.tag_remove("italic", tk.SEL_FIRST, tk.SEL_LAST)
            else:
                self.text_area.tag_add("italic", tk.SEL_FIRST, tk.SEL_LAST)
            self.log_message("Toggled italic on selected text.")
        except tk.TclError:
            self.log_message("No text selected to apply italic.")

    def search_text(self, word):
        """Search for a specific word in the text area."""
        self.text_area.tag_remove("search", 1.0, tk.END)
        if word:
            start_pos = 1.0
            while True:
                start_pos = self.text_area.search(word, start_pos, stopindex=tk.END)
                if not start_pos:
                    break
                end_pos = f"{start_pos}+{len(word)}c"
                self.text_area.tag_add("search", start_pos, end_pos)
                start_pos = end_pos
            self.text_area.tag_config("search", background="yellow")
            self.log_message(f"Search completed for: {word}")

    def replace_text(self, word, replacement):
        """Replace a specific word in the text area."""
        if word and replacement:
            content = self.text_area.get(1.0, tk.END)
            new_content = content.replace(word, replacement)
            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, new_content)
            self.log_message(f"Replaced '{word}' with '{replacement}'.")

if __name__ == "__main__":
    root = tk.Tk()
    app = AdvancedNotepadApp(root)
    root.mainloop()
